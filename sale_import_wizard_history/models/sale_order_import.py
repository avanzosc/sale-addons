# Copyright 2026 Eñaut Alberdi - AvanzOSC
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import unicodedata
from io import BytesIO

from odoo import _, fields, models
from odoo.exceptions import UserError
from odoo.models import expression
from odoo.tools.safe_eval import safe_eval

from odoo.addons.base_import_wizard.models.base_import import check_number, convert2date

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class SaleOrderImport(models.Model):
    _inherit = "sale.order.import"

    historical = fields.Boolean(
        copy=False,
        index=True,
    )

    def _check_import_line_schema(self):
        line_model = self.env["sale.order.import.line"]
        self.env.cr.execute(
            """
            SELECT column_name
              FROM information_schema.columns
             WHERE table_schema = current_schema()
               AND table_name = %s
            """,
            [line_model._table],
        )
        existing_columns = {column for (column,) in self.env.cr.fetchall()}
        required_columns = {"included_tax_amount_import"}
        missing_columns = sorted(required_columns - existing_columns)
        if missing_columns:
            raise UserError(
                _(
                    "The import cannot continue because the database schema is "
                    "outdated. "
                    "Missing columns on table '%(table)s': %(columns)s. "
                    "Please upgrade module 'sale_import_wizard_history' and try again."
                )
                % {
                    "table": line_model._table,
                    "columns": ", ".join(missing_columns),
                }
            )

    def action_import_file(self):
        self._check_import_line_schema()
        return super().action_import_file()

    @staticmethod
    def _normalize_header(value):
        value = "".join(
            ch
            for ch in unicodedata.normalize("NFD", str(value or "").lower())
            if unicodedata.category(ch) != "Mn"
        )
        return "".join(ch for ch in value if ch.isalnum())

    @staticmethod
    def _get_column_value(row_values, column):
        value = row_values.get(column)
        if value not in (None, ""):
            return value
        normalized_column = SaleOrderImport._normalize_header(column)
        for key, key_value in row_values.items():
            normalized_key = SaleOrderImport._normalize_header(key)
            if normalized_key == normalized_column:
                return key_value
        return ""

    @staticmethod
    def _to_float(value):
        if value in (None, ""):
            return 0.0
        number = check_number(value)
        if number is False:
            return 0.0
        return float(number)

    @staticmethod
    def _to_bool(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, int | float):
            return bool(value)
        text = str(value or "").strip().lower()
        return text in {"1", "true", "t", "yes", "y", "x"}

    def _get_line_fields_values(self, row_values):
        values = super()._get_line_fields_values(row_values)
        get_value = self._get_column_value

        def first_value(*columns):
            for column in columns:
                value = get_value(row_values, column)
                if value not in (None, ""):
                    return value
            return None

        def string_value(current, *columns):
            value = first_value(*columns)
            if value in (None, ""):
                return current or ""
            return str(value).strip()

        def float_value(current, *columns):
            value = first_value(*columns)
            if value in (None, ""):
                return current
            return self._to_float(value)

        def bool_value(current, *columns):
            value = first_value(*columns)
            if value in (None, ""):
                return current
            return self._to_bool(value)

        confirmation_date = first_value("Confirmation Date", "Fecha de confirmación")
        date_order = first_value("Order Date", "FechaPedido")
        delivery_date = first_value("Expected Shipment Date", "FechaEntrega")
        quantity_ordered = first_value("QuantityOrdered", "Cantidad")
        item_price = first_value("Item Price", "PrecioUnitario")
        total_order_amount = first_value(
            "TotalImportePedido",
            "Total de pedido",
            "Total Pedido",
        )
        included_tax_raw = first_value("Is Inclusive Tax", "Impuesto incluido")
        included_tax_amount = values.get("included_tax_amount_import")
        if included_tax_raw not in (None, ""):
            included_tax_amount = self._to_float(included_tax_raw)
        converted_confirmation_date = (
            convert2date(confirmation_date)
            if confirmation_date not in (None, "")
            else False
        )
        values.update(
            {
                "salesorder_external_id": string_value(
                    values.get("salesorder_external_id"),
                    "SalesOrder ID",
                    "Sales Order ID",
                    "Pedido de Venta",
                    "Origin",
                ),
                "origin_import": string_value(values.get("origin_import"), "Origin"),
                "sales_team_import": string_value(
                    values.get("sales_team_import"),
                    "Equipo de ventas (team_id)",
                    "Equipo de ventas(team_id)",
                    "Equipo de ventas",
                    "Sales Channel",
                ),
                "currency_code_import": string_value(
                    values.get("currency_code_import"),
                    "Currency Code",
                    "Código de Moneda",
                ),
                "entity_discount_percent": float_value(
                    values.get("entity_discount_percent"), "Entity Discount Percent"
                ),
                "client_order_ref": string_value(
                    values.get("client_order_ref"),
                    "SalesOrder Number",
                    "NumeroPedidoCliente",
                ),
                "product_name": string_value(
                    values.get("product_name"), "Item Name", "NombreProducto"
                ),
                "product_code": string_value(
                    values.get("product_code"), "SKU", "CodigoProducto"
                ),
                "product_barcode": string_value(
                    values.get("product_barcode"), "EAN", "CodigoBarrasProducto"
                ),
                "customer_name": string_value(
                    values.get("customer_name"), "Customer Name", "NombreCliente"
                ),
                "customer_reference": string_value(
                    values.get("customer_reference"), "Customer ID", "ReferenciaCliente"
                ),
                "order_name": string_value(
                    values.get("order_name"),
                    "Origin",
                    "SalesOrder Number",
                    "NumeroPedidoCliente",
                    "Pedido de Venta",
                ),
                "confirmation_date": (
                    converted_confirmation_date
                    if confirmation_date not in (None, "")
                    else values.get("confirmation_date")
                ),
                "date_order": (
                    convert2date(date_order) if date_order else values.get("date_order")
                ),
                "delivery_date": (
                    convert2date(delivery_date)
                    if delivery_date
                    else values.get("delivery_date")
                ),
                "discount": float_value(values.get("discount"), "Discount"),
                "product_uos_qty": float_value(
                    values.get("product_uos_qty"), "QuantityOrdered", "Cantidad"
                ),
                "quantity_invoiced": float_value(
                    values.get("quantity_invoiced"), "QuantityInvoiced"
                ),
                "quantity_packed": float_value(
                    values.get("quantity_packed"), "QuantityPacked"
                ),
                "product_uom_name": string_value(
                    values.get("product_uom_name"),
                    "Usage unit (UOM)",
                    "Usage Unit(UOM)",
                    "Usage unit",
                ),
                "usage_unit": string_value(
                    values.get("usage_unit"),
                    "Usage unit (UOM)",
                    "Usage Unit(UOM)",
                    "Usage unit",
                ),
                "warehouse_name_import": string_value(
                    values.get("warehouse_name_import"), "Warehouse Name"
                ),
                "line_type": string_value(
                    values.get("line_type"), "LineType", "Status", "Estado"
                ),
                "salesman_name": string_value(
                    values.get("salesman_name"), "Sales Person"
                ),
                "invoiced": bool_value(values.get("invoiced"), "Invoiced"),
                "quantity": (
                    self._to_float(quantity_ordered)
                    if quantity_ordered not in (None, "")
                    else values.get("quantity")
                ),
                "price_unit": (
                    self._to_float(item_price)
                    if item_price not in (None, "")
                    else values.get("price_unit")
                ),
                "total_order_amount": (
                    self._to_float(total_order_amount)
                    if self.historical and total_order_amount not in (None, "")
                    else values.get("total_order_amount")
                ),
                "line_subtotal_amount": float_value(
                    values.get("line_subtotal_amount"),
                    "SubTotal",
                    "Subtotal de línea",
                    "Subtotal de linea",
                ),
                "item_total_amount": float_value(
                    values.get("item_total_amount"),
                    "Item Total",
                    "Total de artículo",
                    "Total de articulo",
                ),
                "line_total_amount": float_value(
                    values.get("line_total_amount"),
                    "Line Total",
                    "Total",
                    "Total de línea",
                    "Total de linea",
                ),
                "payment_terms_import": string_value(
                    values.get("payment_terms_import"), "Payment Terms"
                ),
                "payment_terms_label_import": string_value(
                    values.get("payment_terms_label_import"), "Payment Terms Label"
                ),
                "purchase_order_import": string_value(
                    values.get("purchase_order_import"), "PurchaseOrder"
                ),
                "exchange_rate_import": float_value(
                    values.get("exchange_rate_import"),
                    "Exchange Rate",
                    "Tipo de cambio",
                ),
                "is_inclusive_tax_import": (
                    bool(included_tax_amount)
                    if included_tax_raw not in (None, "")
                    else values.get("is_inclusive_tax_import")
                ),
                "included_tax_amount_import": included_tax_amount,
            }
        )
        return values

    def button_open_import_line(self):
        self.ensure_one()
        action = super().button_open_import_line()
        base_domain = action.get("domain") or []
        action["domain"] = expression.AND(
            [base_domain, [("historical", "=", self.historical)]]
        )
        context = action.get("context") or {}
        if isinstance(context, str):
            context = safe_eval(context)
        context.update(
            {
                "search_default_historical_true": 1 if self.historical else 0,
                "search_default_historical_false": 0 if self.historical else 1,
            }
        )
        if self.historical:
            action["view_mode"] = "list,pivot,form"
        action["context"] = context
        return action

    def action_process(self):
        if self.filtered("historical"):
            raise UserError(_("Historical imports cannot be processed."))
        return super().action_process()

    def _read_xlsx(self):
        if not load_workbook:
            raise UserError(
                _('Unable to load "xlsx" file: requires Python module "openpyxl"')
            )
        lines = []
        workbook = load_workbook(
            filename=BytesIO(base64.decodebytes(self.data)),
            read_only=True,
            data_only=True,
        )
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            keys = next(rows, None)
            if not keys:
                continue
            keys = ["" if key is None else str(key).strip() for key in keys]
            for row in rows:
                values = dict(zip(keys, row, strict=False))
                line_data = self._get_line_values(values)
                if line_data:
                    lines.append((0, 0, line_data))
        return lines
